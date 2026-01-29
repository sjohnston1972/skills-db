#!/usr/bin/env python3
"""
Import Real Skills Matrix data from Excel file
"""

import openpyxl

# Map text skill levels to numeric (0-5)
SKILL_LEVEL_MAP = {
    'no exposure': 0,
    'shadowed or lab deployed': 1,
    'training only': 1,
    'implemented once': 2,
    'implemented multiple time': 3,
    'implemented multiple times': 3,
    'subject matter expert': 5,
}

def normalize_skill_level(text):
    """Convert text skill level to numeric 0-5"""
    if not text:
        return 0
    text_lower = str(text).strip().lower()
    return SKILL_LEVEL_MAP.get(text_lower, 0)

def parse_excel(filepath):
    """Parse the Excel file"""
    print(f"Reading: {filepath}\n")
    wb = openpyxl.load_workbook(filepath, data_only=True)

    # Skip certain sheets
    skip_sheets = {'certs', 'learning paths', 'offshore and clearance'}

    data = {
        'main_skills': {},
        'resources': {}
    }

    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in skip_sheets:
            continue

        sheet = wb[sheet_name]
        print(f"Processing: {sheet_name}")

        # Row 1 is header: Engineer, SubSkill1, SubSkill2, ...
        header_row = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]

        if not header_row or not header_row[0]:
            print(f"  Skipped - no header\n")
            continue

        # Extract sub-skill names (columns 2+)
        sub_skills = []
        for col_idx, cell_value in enumerate(header_row[1:], start=2):
            if cell_value and str(cell_value).strip():
                sub_skill = str(cell_value).strip()
                sub_skills.append(sub_skill)

        if not sub_skills:
            print(f"  Skipped - no sub-skills\n")
            continue

        print(f"  Sub-skills: {len(sub_skills)}")

        # Store main skill
        data['main_skills'][sheet_name] = {
            'name': sheet_name,
            'sub_skills': sub_skills
        }

        # Parse engineers and their skill levels (rows 2+)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            engineer_name = row[0]

            if not engineer_name or str(engineer_name).strip() == '':
                continue

            engineer_name = str(engineer_name).strip()

            # Initialize engineer if not exists
            if engineer_name not in data['resources']:
                data['resources'][engineer_name] = {
                    'name': engineer_name,
                    'email': f"{engineer_name.lower().replace(' ', '.')}@company.com",
                    'skills': {}
                }

            # Extract skill levels for this engineer
            data['resources'][engineer_name]['skills'][sheet_name] = {}

            for idx, sub_skill in enumerate(sub_skills):
                col_idx = idx + 1  # Column index in row (0=engineer, 1+=skills)
                if col_idx < len(row):
                    skill_text = row[col_idx]
                    skill_level = normalize_skill_level(skill_text)
                    if skill_level > 0:  # Only store non-zero skills
                        data['resources'][engineer_name]['skills'][sheet_name][sub_skill] = skill_level

        print(f"  Engineers: {len([r for r in data['resources'].values() if sheet_name in r['skills']])}\n")

    return data

def generate_sql(data):
    """Generate SQL"""
    lines = []

    # Clear existing
    lines.append("-- Clear existing data")
    lines.append("TRUNCATE TABLE resource_sub_skills, sub_skills, main_skills, resources CASCADE;")
    lines.append("")

    # Main skills
    lines.append("-- Main Skills")
    skill_id_map = {}
    for idx, (name, info) in enumerate(sorted(data['main_skills'].items()), start=1):
        skill_id = f"skill-{idx:03d}"
        skill_id_map[name] = skill_id

        # Determine category
        name_lower = name.lower()
        if 'azure' in name_lower or 'aws' in name_lower:
            category = 'Cloud'
        elif 'security' in name_lower:
            category = 'Security'
        elif 'collaboration' in name_lower or 'call' in name_lower:
            category = 'Communication'
        elif 'data' in name_lower or 'nutanix' in name_lower or 'dell' in name_lower:
            category = 'Infrastructure'
        else:
            category = 'Networking'

        # Determine weight (importance)
        weight = 7  # default
        if category == 'Cloud':
            weight = 9
        elif category == 'Security':
            weight = 9
        elif 'cisco' in name_lower and 'sd-wan' in name_lower:
            weight = 8

        name_esc = name.replace("'", "''")
        lines.append(f"INSERT INTO main_skills (id, name, category, weight, skill_type) VALUES ('{skill_id}', '{name_esc}', '{category}', {weight}, 'technical');")

    lines.append("")

    # Sub-skills
    lines.append("-- Sub-Skills")
    for name, info in sorted(data['main_skills'].items()):
        skill_id = skill_id_map[name]
        for sub_skill in info['sub_skills']:
            sub_esc = sub_skill.replace("'", "''")
            lines.append(f"INSERT INTO sub_skills (main_skill_id, name) VALUES ('{skill_id}', '{sub_esc}');")

    lines.append("")

    # Resources
    lines.append("-- Resources")
    resource_id_map = {}
    for idx, (name, info) in enumerate(sorted(data['resources'].items()), start=1):
        resource_id = f"eng-{idx:03d}"
        resource_id_map[name] = resource_id
        name_esc = name.replace("'", "''")
        email_esc = info['email'].replace("'", "''")
        lines.append(f"INSERT INTO resources (id, name, email) VALUES ('{resource_id}', '{name_esc}', '{email_esc}');")

    lines.append("")

    # Skill levels
    lines.append("-- Resource Skill Levels")
    for eng_name, eng_info in sorted(data['resources'].items()):
        resource_id = resource_id_map[eng_name]
        for main_skill, sub_skills in sorted(eng_info['skills'].items()):
            if main_skill in skill_id_map:
                skill_id = skill_id_map[main_skill]
                for sub_skill, level in sorted(sub_skills.items()):
                    sub_esc = sub_skill.replace("'", "''")
                    lines.append(
                        f"INSERT INTO resource_sub_skills (resource_id, sub_skill_id, level) "
                        f"SELECT '{resource_id}', id, {level} FROM sub_skills "
                        f"WHERE main_skill_id = '{skill_id}' AND name = '{sub_esc}';"
                    )

    lines.append("")
    lines.append("-- Update metadata")
    lines.append("UPDATE metadata SET value = CURRENT_TIMESTAMP::TEXT WHERE key = 'last_updated';")

    return '\n'.join(lines)

def main():
    excel_file = '/host-projects/web-projects/skills-db/Project Services Engineers Skills Matrix.xlsx'

    # Parse
    data = parse_excel(excel_file)

    # Summary
    print(f"\n{'='*70}")
    print("SUMMARY")
    print(f"{'='*70}")
    print(f"Main Skills: {len(data['main_skills'])}")
    total_sub_skills = sum(len(s['sub_skills']) for s in data['main_skills'].values())
    print(f"Sub-Skills: {total_sub_skills}")
    print(f"Resources: {len(data['resources'])}")

    total_assignments = 0
    for eng in data['resources'].values():
        for main_skill, sub_skills in eng['skills'].items():
            total_assignments += len(sub_skills)
    print(f"Skill Assignments: {total_assignments}")

    # Generate SQL
    sql = generate_sql(data)

    # Write
    output = '/host-projects/web-projects/skills-db/imported-data.sql'
    with open(output, 'w') as f:
        f.write(sql)

    print(f"\n✅ Generated: {output}")
    print(f"   Lines: {len(sql.splitlines())}")

if __name__ == '__main__':
    main()
